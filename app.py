import streamlit as st
import pandas as pd
import sqlite3
import io
import time
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# --- KONFIGURASI DATABASE ---
DB_NAME = "pesantren.db"

def init_db():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS profil_pesantren (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nama_pp TEXT,
            nspp TEXT,
            alamat TEXT,
            pengasuh TEXT,
            akta_notaris TEXT,
            sk_ijin TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS satuan_pendidikan (
            npsn TEXT PRIMARY KEY,
            nama_satuan TEXT NOT NULL,
            jenjang TEXT,
            status TEXT,
            kepala_sekolah TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS ustad (
            id_ustad TEXT PRIMARY KEY,
            nama TEXT NOT NULL,
            bidang_pengajaran TEXT,
            no_hp TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS santri (
            nis TEXT PRIMARY KEY,
            nama TEXT NOT NULL,
            kelas TEXT,
            kamar TEXT,
            wali_santri TEXT
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS sarpras (
            kode_barang TEXT PRIMARY KEY,
            nama_barang TEXT NOT NULL,
            kondisi TEXT,
            jumlah INTEGER
        )
    ''')
    conn.commit()
    conn.close()

init_db()

def run_query(query, params=()):
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute(query, params)
    conn.commit()
    conn.close()

def load_data(table_name):
    conn = sqlite3.connect(DB_NAME)
    df = pd.read_sql_query(f"SELECT * FROM {table_name}", conn)
    conn.close()
    return df

# --- FUNGSI GENERATE WORD ---
def generate_word_portfolio():
    doc = Document()

    df_profil = load_data("profil_pesantren")
    nama_pp = df_profil['nama_pp'].values[0] if not df_profil.empty else "PONDOK PESANTREN"
    
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run(f"PORTOFOLIO & PROFIL LEMBAGA\n{nama_pp.upper()}")
    run_title.bold = True
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = RGBColor(0, 102, 204)

    doc.add_paragraph("-------------------------------------------------------------------------------------------------------------------")

    doc.add_heading("1. Profil & Legalitas Lembaga", level=1)
    if not df_profil.empty:
        p = df_profil.iloc[0]
        doc.add_paragraph(f"Nama Pesantren: {p['nama_pp']}")
        doc.add_paragraph(f"NSPP: {p['nspp'] if p['nspp'] else '-'}")
        doc.add_paragraph(f"Pengasuh / Pimpinan: {p['pengasuh'] if p['pengasuh'] else '-'}")
        doc.add_paragraph(f"Alamat: {p['alamat'] if p['alamat'] else '-'}")
        doc.add_paragraph(f"Akta Notaris: {p['akta_notaris'] if p['akta_notaris'] else '-'}")
        doc.add_paragraph(f"SK Izin Operasional: {p['sk_ijin'] if p['sk_ijin'] else '-'}")
    else:
        doc.add_paragraph("Belum ada data profil lembaga.")

    doc.add_heading("2. Satuan Pendidikan / Unit Sekolah", level=1)
    df_satuan = load_data("satuan_pendidikan")
    if not df_satuan.empty:
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text = 'NPSN', 'Nama Satuan', 'Jenjang', 'Status', 'Kepala Sekolah'
        for _, row in df_satuan.iterrows():
            r = table.add_row().cells
            r[0].text, r[1].text, r[2].text, r[3].text, r[4].text = str(row['npsn']), str(row['nama_satuan']), str(row['jenjang']), str(row['status']), str(row['kepala_sekolah'])
    else:
        doc.add_paragraph("Belum ada data satuan pendidikan.")

    doc.add_heading("3. Data Ustad / Pengajar", level=1)
    df_ustad = load_data("ustad")
    if not df_ustad.empty:
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = 'ID Ustad', 'Nama Lengkap', 'Bidang Pengajaran', 'No. HP'
        for _, row in df_ustad.iterrows():
            r = table.add_row().cells
            r[0].text, r[1].text, r[2].text, r[3].text = str(row['id_ustad']), str(row['nama']), str(row['bidang_pengajaran']), str(row['no_hp'])
    else:
        doc.add_paragraph("Belum ada data ustad.")

    doc.add_heading("4. Data Santri", level=1)
    df_santri = load_data("santri")
    if not df_santri.empty:
        doc.add_paragraph(f"Total Santri Terdaftar: {len(df_santri)} Santri")
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text, hdr[4].text = 'NIS', 'Nama Santri', 'Kelas', 'Kamar', 'Wali Santri'
        for _, row in df_santri.iterrows():
            r = table.add_row().cells
            r[0].text, r[1].text, r[2].text, r[3].text, r[4].text = str(row['nis']), str(row['nama']), str(row['kelas']), str(row['kamar']), str(row['wali_santri'])
    else:
        doc.add_paragraph("Belum ada data santri.")

    doc.add_heading("5. Inventaris Sarana & Prasarana", level=1)
    df_sarpras = load_data("sarpras")
    if not df_sarpras.empty:
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = 'Kode Barang', 'Nama Barang', 'Kondisi', 'Jumlah'
        for _, row in df_sarpras.iterrows():
            r = table.add_row().cells
            r[0].text, r[1].text, r[2].text, r[3].text = str(row['kode_barang']), str(row['nama_barang']), str(row['kondisi']), str(row['jumlah'])
    else:
        doc.add_paragraph("Belum ada data sarpras.")

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

# --- TAMPILAN APLIKASI STREAMLIT ---
st.set_page_config(page_title="SIM Pondok Pesantren", page_icon="🏫", layout="wide")

st.sidebar.title("🏫 SIM Pesantren")
menu = st.sidebar.selectbox("Pilih Menu", [
    "Dashboard", 
    "Profil Pesantren", 
    "Satuan Pendidikan", 
    "Data Ustad", 
    "Data Santri", 
    "Data Sarpras",
    "📄 Cetak Portofolio Word"
])

# 1. DASHBOARD
if menu == "Dashboard":
    conn = sqlite3.connect(DB_NAME)
    df_profil = pd.read_sql_query("SELECT nama_pp FROM profil_pesantren ORDER BY id DESC LIMIT 1", conn)
    conn.close()
    
    nama_pesantren = df_profil['nama_pp'].values[0] if not df_profil.empty else "Pondok Pesantren"
    
    st.header("📊 Dashboard Rekapitulasi")
    st.subheader(f"🏫 {nama_pesantren}")
    st.divider()

    df_satuan = load_data("satuan_pendidikan")
    df_ustad = load_data("ustad")
    df_santri = load_data("santri")
    df_sarpras = load_data("sarpras")

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Satuan Pendidikan", len(df_satuan))
    col2.metric("Total Ustad", len(df_ustad))
    col3.metric("Total Santri", len(df_santri))
    col4.metric("Total Sarpras", len(df_sarpras))

    st.success("✅ Aplikasi siap digunakan dan data tersimpan dengan aman!")

# 2. PROFIL PESANTREN
elif menu == "Profil Pesantren":
    st.header("🏫 Profil & Legalitas Pondok Pesantren")
    conn = sqlite3.connect(DB_NAME)
    df_profil = pd.read_sql_query("SELECT * FROM profil_pesantren ORDER BY id DESC LIMIT 1", conn)
    conn.close()
    
    with st.form("form_profil"):
        st.subheader("Perbarui Informasi Lembaga")
        def_nama = df_profil['nama_pp'].values[0] if not df_profil.empty else ""
        def_nspp = df_profil['nspp'].values[0] if not df_profil.empty else ""
        def_alamat = df_profil['alamat'].values[0] if not df_profil.empty else ""
        def_pengasuh = df_profil['pengasuh'].values[0] if not df_profil.empty else ""
        def_akta = df_profil['akta_notaris'].values[0] if not df_profil.empty else ""
        def_sk = df_profil['sk_ijin'].values[0] if not df_profil.empty else ""
        
        nama_pp = st.text_input("Nama Pondok Pesantren", value=def_nama)
        nspp = st.text_input("Nomor Statistik Pesantren (NSPP)", value=def_nspp)
        alamat = st.text_area("Alamat Lengkap", value=def_alamat)
        pengasuh = st.text_input("Nama Pengasuh / Pimpinan", value=def_pengasuh)
        akta_notaris = st.text_input("Nomor & Tanggal Akta Notaris", value=def_akta, placeholder="Contoh: Akta No. 12 / 15 Januari 2020")
        sk_ijin = st.text_input("Nomor & Tanggal SK Ijin Operasional", value=def_sk, placeholder="Contoh: SK.123/Kemenag/2021 / 10 Mei 2021")
        
        if st.form_submit_button("Simpan / Perbarui Profil"):
            if nama_pp:
                run_query("DELETE FROM profil_pesantren")
                run_query("INSERT INTO profil_pesantren (nama_pp, nspp, alamat, pengasuh, akta_notaris, sk_ijin) VALUES (?, ?, ?, ?, ?, ?)", 
                          (nama_pp, nspp, alamat, pengasuh, akta_notaris, sk_ijin))
                st.success("✅ Profil pesantren berhasil diperbarui!")
                st.rerun()
            else:
                st.warning("Nama Pondok Pesantren wajib diisi.")

    if not df_profil.empty:
        st.divider()
        st.subheader("📄 Informasi Resmi Lembaga")
        p = df_profil.iloc[0]
        col1, col2 = st.columns(2)
        with col1:
            st.markdown(f"**Nama Pondok Pesantren:**\n{p['nama_pp']}")
            st.markdown(f"**Nomor Statistik (NSPP):**\n{p['nspp'] if p['nspp'] else '-'}")
            st.markdown(f"**Pengasuh / Pimpinan:**\n{p['pengasuh'] if p['pengasuh'] else '-'}")
        with col2:
            st.markdown(f"**Alamat:**\n{p['alamat'] if p['alamat'] else '-'}")
            st.markdown(f"**Akta Notaris:**\n{p['akta_notaris'] if p['akta_notaris'] else '-'}")
            st.markdown(f"**SK Ijin Operasional:**\n{p['sk_ijin'] if p['sk_ijin'] else '-'}")

# 3. SATUAN PENDIDIKAN
elif menu == "Satuan Pendidikan":
    st.header("🏫 Manajemen Satuan Pendidikan / Unit Sekolah")
    tab1, tab2 = st.tabs(["➕ Input Manual", "📁 Impor dari Excel"])

    with tab1:
        with st.form("form_satuan", clear_on_submit=True):
            npsn = st.text_input("NPSN / Nomor Induk Satuan (Kosongkan jika non-formal)")
            nama_satuan = st.text_input("Nama Satuan Pendidikan (Contoh: MTs Darul Huda)")
            jenjang = st.selectbox("Jenjang Pendidikan", ["RA / TK", "MI / SD", "MTs / SMP", "MA / SMA / SMK", "Madrasah Diniyah", "TPQ / TKA", "Lainnya"])
            status = st.selectbox("Status Lembaga", ["Swasta", "Negeri"])
            kepala_sekolah = st.text_input("Nama Kepala / Pimpinan Unit")
            
            if st.form_submit_button("Simpan Satuan Pendidikan"):
                if nama_satuan:
                    if not npsn:
                        npsn = f"NON-{int(time.time())}"
                    try:
                        run_query("INSERT INTO satuan_pendidikan VALUES (?, ?, ?, ?, ?)", (npsn, nama_satuan, jenjang, status, kepala_sekolah))
                        st.success(f"Satuan pendidikan **{nama_satuan}** berhasil disimpan!")
                        st.rerun()
                    except sqlite3.IntegrityError:
                        st.error(f"NPSN '{npsn}' sudah terdaftar.")

    with tab2:
        st.subheader("Unggah Data Satuan Pendidikan via Excel (.xlsx)")
        df_template_satuan = pd.DataFrame(columns=["npsn", "nama_satuan", "jenjang", "status", "kepala_sekolah"])
        output_satuan = io.BytesIO()
        with pd.ExcelWriter(output_satuan, engine='openpyxl') as writer:
            df_template_satuan.to_excel(writer, sheet_name='Satuan Pendidikan', index=False)
        st.download_button("📥 Unduh Template Excel", data=output_satuan.getvalue(), file_name="template_satuan_pendidikan.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        
        st.divider()
        uploaded_file_satuan = st.file_uploader("Pilih file Excel (.xlsx)", type=["xlsx"], key="up_satuan")
        if uploaded_file_satuan is not None:
            try:
                df_upload_satuan = pd.read_excel(uploaded_file_satuan, engine='openpyxl')
                if st.button("Proses Impor Satuan Pendidikan"):
                    conn = sqlite3.connect(DB_NAME)
                    df_upload_satuan.to_sql('satuan_pendidikan', conn, if_exists='append', index=False)
                    conn.close()
                    st.success("✅ Data satuan pendidikan berhasil diimpor!")
                    st.rerun()
            except Exception as e:
                st.error(f"Terjadi kesalahan: {e}")

    st.divider()
    st.subheader("📋 Daftar Unit Satuan Pendidikan")
    st.dataframe(load_data("satuan_pendidikan"), use_container_width=True)

# 4. DATA USTAD
elif menu == "Data Ustad":
    st.header("👨‍🏫 Manajemen Data Ustad / Pengajar")
    tab1, tab2, tab3 = st.tabs(["➕ Input Manual", "📁 Impor dari Excel", "✏️ Edit / Hapus Data"])
    
    with tab1:
        with st.form("form_ustad", clear_on_submit=True):
            id_ustad = st.text_input("ID Ustad")
            nama = st.text_input("Nama Lengkap & Gelar")
            bidang = st.text_input("Bidang Pengajaran / Kitab")
            hp = st.text_input("No. HP / WhatsApp")
            if st.form_submit_button("Simpan Data Ustad"):
                if id_ustad and nama:
                    try:
                        run_query("INSERT INTO ustad VALUES (?, ?, ?, ?)", (id_ustad, nama, bidang, hp))
                        st.success(f"Data ustad **{nama}** berhasil disimpan!")
                        st.rerun()
                    except sqlite3.IntegrityError:
                        st.error(f"ID Ustad '{id_ustad}' sudah terdaftar.")

    with tab2:
        st.subheader("Unggah Data Ustad Masal via Excel (.xlsx)")
        df_template_ustad = pd.DataFrame(columns=["id_ustad", "nama", "bidang_pengajaran", "no_hp"])
        output_ustad = io.BytesIO()
        with pd.ExcelWriter(output_ustad, engine='openpyxl') as writer:
            df_template_ustad.to_excel(writer, sheet_name='Data Ustad', index=False)
        st.download_button("📥 Unduh Template Excel", data=output_ustad.getvalue(), file_name="template_data_ustad.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        
        st.divider()
        uploaded_file_ustad = st.file_uploader("Pilih file Excel (.xlsx) data ustad", type=["xlsx"], key="up_ustad")
        if uploaded_file_ustad is not None:
            try:
                df_upload_ustad = pd.read_excel(uploaded_file_ustad, engine='openpyxl')
                if st.button("Proses Impor Ustad"):
                    conn = sqlite3.connect(DB_NAME)
                    df_upload_ustad.to_sql('ustad', conn, if_exists='append', index=False)
                    conn.close()
                    st.success("✅ Data ustad berhasil diimpor!")
                    st.rerun()
            except Exception as e:
                st.error(f"Terjadi kesalahan: {e}")

    with tab3:
        st.subheader("Pilih Ustad yang Ingin Diedit atau Dihapus")
        df_ustad_s = load_data("ustad")
        if not df_ustad_s.empty:
            pilih_ustad = st.selectbox("Pilih ID & Nama Ustad", options=df_ustad_s['id_ustad'].tolist(), format_func=lambda x: f"{x} - {df_ustad_s.loc[df_ustad_s['id_ustad'] == x, 'nama'].values[0]}")
            
            data_terpilih = df_ustad_s[df_ustad_s['id_ustad'] == pilih_ustad].iloc[0]
            
            with st.form("form_edit_ustad"):
                new_nama = st.text_input("Nama Lengkap & Gelar", value=data_terpilih['nama'])
                new_bidang = st.text_input("Bidang Pengajaran", value=data_terpilih['bidang_pengajaran'])
                new_hp = st.text_input("No. HP / WhatsApp", value=data_terpilih['no_hp'])
                
                col_e1, col_e2 = st.columns(2)
                update_btn = col_e1.form_submit_button("💾 Perbarui Data")
                delete_btn = col_e2.form_submit_button("🗑️ Hapus Data", type="primary")
                
                if update_btn:
                    run_query("UPDATE ustad SET nama = ?, bidang_pengajaran = ?, no_hp = ? WHERE id_ustad = ?", (new_nama, new_bidang, new_hp, pilih_ustad))
                    st.success("✅ Data ustad berhasil diperbarui!")
                    st.rerun()
                elif delete_btn:
                    run_query("DELETE FROM ustad WHERE id_ustad = ?", (pilih_ustad,))
                    st.success("🗑️ Data ustad berhasil dihapus!")
                    st.rerun()
        else:
            st.info("Belum ada data ustad.")

    st.divider()
    st.subheader("Daftar Ustad Aktif")
    st.dataframe(load_data("ustad"), use_container_width=True)

# 5. DATA SANTRI
elif menu == "Data Santri":
    st.header("👦👧 Manajemen Data Santri")
    tab1, tab2, tab3 = st.tabs(["➕ Input Manual", "📁 Impor dari Excel", "✏️ Edit / Hapus Data"])
    
    with tab1:
        with st.form("form_santri", clear_on_submit=True):
            nis = st.text_input("NIS")
            nama = st.text_input("Nama Lengkap")
            kelas = st.text_input("Kelas")
            kamar = st.text_input("Kamar Asrama")
            wali = st.text_input("Nama Wali")
            if st.form_submit_button("Simpan Data Santri"):
                if nis and nama:
                    try:
                        run_query("INSERT INTO santri VALUES (?, ?, ?, ?, ?)", (nis, nama, kelas, kamar, wali))
                        st.success(f"Data santri **{nama}** berhasil disimpan!")
                        st.rerun()
                    except sqlite3.IntegrityError:
                        st.error(f"NIS '{nis}' sudah terdaftar.")

    with tab2:
        st.subheader("Unggah Data Santri Masal via Excel (.xlsx)")
        df_template_santri = pd.DataFrame(columns=["nis", "nama", "kelas", "kamar", "wali_santri"])
        output_santri = io.BytesIO()
        with pd.ExcelWriter(output_santri, engine='openpyxl') as writer:
            df_template_santri.to_excel(writer, sheet_name='Data Santri', index=False)
        st.download_button("📥 Unduh Template Excel", data=output_santri.getvalue(), file_name="template_data_santri.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        
        st.divider()
        uploaded_file = st.file_uploader("Pilih file Excel (.xlsx) data santri", type=["xlsx"], key="up_santri")
        if uploaded_file is not None:
            try:
                df_upload = pd.read_excel(uploaded_file, engine='openpyxl')
                if st.button("Proses Impor Santri"):
                    conn = sqlite3.connect(DB_NAME)
                    df_upload.to_sql('santri', conn, if_exists='append', index=False)
                    conn.close()
                    st.success("✅ Data santri berhasil diimpor!")
                    st.rerun()
            except Exception as e:
                st.error(f"Terjadi kesalahan: {e}")

    with tab3:
        st.subheader("Pilih Santri yang Ingin Diedit atau Dihapus")
        df_santri_s = load_data("santri")
        if not df_santri_s.empty:
            pilih_santri = st.selectbox("Pilih NIS & Nama Santri", options=df_santri_s['nis'].tolist(), format_func=lambda x: f"{x} - {df_santri_s.loc[df_santri_s['nis'] == x, 'nama'].values[0]}")
            
            santri_terpilih = df_santri_s[df_santri_s['nis'] == pilih_santri].iloc[0]
            
            with st.form("form_edit_santri"):
                new_nama_s = st.text_input("Nama Lengkap", value=santri_terpilih['nama'])
                new_kelas = st.text_input("Kelas", value=santri_terpilih['kelas'])
                new_kamar = st.text_input("Kamar Asrama", value=santri_terpilih['kamar'])
                new_wali = st.text_input("Nama Wali Santri", value=santri_terpilih['wali_santri'])
                
                col_s1, col_s2 = st.columns(2)
                update_s_btn = col_s1.form_submit_button("💾 Perbarui Data")
                delete_s_btn = col_s2.form_submit_button("🗑️ Hapus Data", type="primary")
                
                if update_s_btn:
                    run_query("UPDATE santri SET nama = ?, kelas = ?, kamar = ?, wali_santri = ? WHERE nis = ?", (new_nama_s, new_kelas, new_kamar, new_wali, pilih_santri))
                    st.success("✅ Data santri berhasil diperbarui!")
                    st.rerun()
                elif delete_s_btn:
                    run_query("DELETE FROM santri WHERE nis = ?", (pilih_santri,))
                    st.success("🗑️ Data santri berhasil dihapus!")
                    st.rerun()
        else:
            st.info("Belum ada data santri.")

    st.divider()
    st.subheader("Daftar Santri Aktif")
    st.dataframe(load_data("santri"), use_container_width=True)

# 6. DATA SARPRAS
elif menu == "Data Sarpras":
    st.header("🏢 Manajemen Sarana & Prasarana")
    with st.form("form_sarpras", clear_on_submit=True):
        kode = st.text_input("Kode Barang")
        nama_brg = st.text_input("Nama Barang")
        kondisi = st.selectbox("Kondisi", ["Baik", "Rusak Ringan", "Rusak Berat"])
        jumlah = st.number_input("Jumlah Unit", min_value=1, step=1)
        if st.form_submit_button("Simpan Data Sarpras"):
            if kode and nama_brg:
                try:
                    run_query("INSERT INTO sarpras VALUES (?, ?, ?, ?)", (kode, nama_brg, kondisi, int(jumlah)))
                    st.success(f"Data sarpras **{nama_brg}** berhasil disimpan!")
                    st.rerun()
                except sqlite3.IntegrityError:
                    st.error(f"Kode Barang '{kode}' sudah terdaftar.")

    st.subheader("Daftar Inventaris Sarpras")
    st.dataframe(load_data("sarpras"), use_container_width=True)

# 7. CETAK PORTOFOLIO WORD
elif menu == "📄 Cetak Portofolio Word":
    st.header("📄 Cetak Portofolio Lembaga (.docx)")
    st.write("Unduh rangkuman seluruh data pesantren per lembaga dalam bentuk dokumen Word yang rapi dan siap dicetak.")
    st.divider()
    
    df_profil = load_data("profil_pesantren")
    nama_pp = df_profil['nama_pp'].values[0] if not df_profil.empty else "Pesantren"
    
    doc_file = generate_word_portfolio()
    
    st.download_button(
        label="📥 Download Portofolio Word (.docx)",
        data=doc_file,
        file_name=f"Portofolio_{nama_pp.replace(' ', '_')}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )
